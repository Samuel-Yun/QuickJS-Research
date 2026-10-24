"""Read-only checks of research sources and PPTX contents; writes a report only."""
from collections import defaultdict
import csv
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
REPORT = ROOT.parent
PROJECT = REPORT.parents[1]
SVG = ROOT / "weekly_quickjs_20261017" / "svg_output"
NS = {"s": "http://www.w3.org/2000/svg"}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraphs(page, group):
    tree = ET.parse(SVG / page)
    found = tree.find(f"s:g[@id='{group}']/s:text", NS)
    assert found is not None
    return [s.strip() for s in found.itertext() if s.strip()]


def main():
    validation = json.loads((REPORT / "validation.json").read_text(encoding="utf-8"))
    changed = [p for p, sha in validation["source_hashes"].items() if digest(PROJECT / p) != sha]
    assert not changed, changed
    data = {"status": "PASS", "unchanged_prior_sources": len(validation["source_hashes"]), "changed_sources": changed}
    data["beamer_main_sha256"] = digest(REPORT / "main.tex")
    data["beamer_zip_sha256"] = digest(REPORT / "weekly_20261017_overleaf.zip")
    assert data["beamer_zip_sha256"] == "54803cbc94e750b9431dc9a13ca6b8558efbdf09c86ea0c53dbb512afe0c84fb"
    s3 = PROJECT / "experiments/richards_site_attribution/attempt01"
    for engine, filename, page in [
        ("quickjs", "TaskControlBlock.prototype.isHeldOrSuspended", "07_qjs_full.svg"),
        ("v8", "TaskControlBlock.isHeldOrSuspended", "08_v8_full.svg"),
    ]:
        decoded = json.loads((s3 / "bytecode" / engine / (filename + ".decoded.json")).read_text(encoding="utf-8"))
        actual = paragraphs(page, "complete-bytecode")
        expected = [f"{i['offset']} {i['opcode']} {i['operands']}".strip() for i in decoded]
        assert actual == expected, (engine, actual, expected)
    data["full_bytecode_exact_match"] = {"quickjs": 16, "v8": 14}
    rows = list(csv.DictReader((s3 / "summary/site_samples.csv").open(encoding="utf-8")))
    samples = defaultdict(int)
    for row in rows:
        samples[row["property"]] += int(row["lookup_samples"])
    assert sum(samples.values()) == 1229
    assert {p: samples[p] for p in ("currentTcb", "state", "run")} == {"currentTcb": 280, "state": 204, "run": 79}
    data["lookup_samples"] = {p: samples[p] for p in ("currentTcb", "state", "run")}
    prs = Presentation(ROOT / "weekly_20261017_pptmaster.pptx")
    assert len(prs.slides) == 8
    tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 2
    timing_table = tables[0]
    s2 = PROJECT / "experiments/call_path_followup/attempt01/summary.csv"
    timings = list(csv.DictReader(s2.open(encoding="utf-8")))
    for r, name in enumerate(("inline", "direct_call", "own_method", "prototype_method"), 1):
        q = next(x for x in timings if x["case"] == name and x["engine"] == "quickjs")
        v = next(x for x in timings if x["case"] == name and x["engine"] == "v8")
        assert timing_table.cell(r, 2).text == f"{float(q['median_ms'])*1000:.3f}"
        assert timing_table.cell(r, 3).text == f"{float(v['median_ms'])*1000:.3f}"
        assert timing_table.cell(r, 4).text == f"{float(q['v8_over_quickjs']):.3f}"
    with zipfile.ZipFile(ROOT / "weekly_20261017_pptmaster.pptx") as z:
        slides = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        pictures = sum(z.read(n).count(b"<p:pic>") for n in slides)
        notes = [n for n in z.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]
        assert len(notes) == 8 and pictures == 0
        data["editable_package"] = {"slides": len(slides), "native_tables": len(tables), "picture_objects": pictures, "notes": len(notes)}
    render = json.loads((ROOT / "preview/powerpoint_render.json").read_text(encoding="utf-8"))
    assert render["slide_count"] == 8 and render["off_slide_text_count"] == 0
    data["powerpoint"] = {"rendered_pages": 8, "off_slide_text": 0, "visually_reviewed_pages": 8}
    data["beamer_compiled"] = False
    data["artifacts"] = {
        str(p.relative_to(ROOT)): {"sha256": digest(p), "bytes": p.stat().st_size}
        for p in (ROOT / "weekly_20261017_pptmaster.pptx", ROOT / "preview/weekly_20261017_pptmaster.pdf")
    }
    (ROOT / "delivery_validation.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(data, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
